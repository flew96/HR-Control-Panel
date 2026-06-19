from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from .models import (
    AuditLog,
    Campaign,
    CampaignRecord,
    Department,
    Employee,
    EmployeeColumnPreference,
    Position,
    UserDepartmentAccess,
)


def make_user(username, hr_perms=(), auth_perms=()):
    User = get_user_model()
    user = User.objects.create_user(username=username, password="Test12345!")
    permissions = list(
        Permission.objects.filter(
            content_type__app_label="hr",
            codename__in=hr_perms,
        )
    )
    if auth_perms:
        permissions += list(
            Permission.objects.filter(
                content_type__app_label="auth",
                content_type__model="user",
                codename__in=auth_perms,
            )
        )
    user.user_permissions.set(permissions)
    return user


class HrFeatureTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.department = Department.objects.create(name="Engineering")
        cls.position = Position.objects.create(name="Backend Developer")
        cls.employee = Employee.objects.create(
            first_name="Иван",
            last_name="Тестов",
            email="ivan@example.test",
            department=cls.department,
            position=cls.position,
            hire_date="2026-01-10",
            is_active=True,
            ipr_status=Employee.IPR_ACTUAL,
            contract_type=Employee.CONTRACT_LABOR,
            contract_end_date="2026-12-31",
            salary_gross="150000",
            salary_net="130500",
            currency="RUB",
        )
        for index in range(30):
            Employee.objects.create(
                first_name=f"User{index}",
                last_name="Page",
                email=f"page{index}@example.test",
                is_active=True,
                ipr_status=Employee.IPR_NONE,
                currency="RUB",
            )

        cls.viewer = make_user("viewer", hr_perms=["view_employee"])
        cls.manager = make_user(
            "manager",
            hr_perms=[
                "view_employee",
                "add_employee",
                "change_employee",
                "view_employeecolumnpreference",
            ],
        )
        cls.directory_admin = make_user(
            "directory_admin",
            hr_perms=[
                "view_position",
                "add_position",
                "change_position",
                "delete_position",
                "view_campaign",
                "view_employee",
                "add_campaignrecord",
                "change_campaignrecord",
            ],
        )

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get("/employees/")

        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response["Location"])

    def test_viewer_cannot_create_employee(self):
        self.client.force_login(self.viewer)

        response = self.client.post(
            "/employees/new/",
            {
                "last_name": "No",
                "first_name": "Rights",
                "ipr_status": Employee.IPR_NONE,
                "currency": "RUB",
            },
        )

        self.assertEqual(response.status_code, 403)

    def test_employee_list_uses_saved_columns_and_pagination(self):
        self.client.force_login(self.viewer)

        response = self.client.post(
            "/employees/columns/",
            {
                "columns": ["full_name", "email"],
                "page_size": "10",
                "next": "/employees/",
            },
        )
        self.assertEqual(response.status_code, 302)
        preference = EmployeeColumnPreference.objects.get(user=self.viewer)
        self.assertEqual(preference.columns, ["full_name", "email"])
        self.assertEqual(preference.page_size, 10)

        response = self.client.get("/employees/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["employees"]), 10)
        self.assertEqual(response.context["visible_columns"], ["full_name", "email"])

    def test_employee_export_returns_xlsx_and_audit_log(self):
        self.client.force_login(self.viewer)

        response = self.client.get("/employees/export/?q=ivan")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        self.assertEqual(worksheet["A1"].value, "ID")
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.ACTION_EXPORT).exists())

    def test_employee_import_creates_employee_and_audit_log(self):
        self.client.force_login(self.manager)
        upload = self._build_import_file(
            [
                [
                    None,
                    "Импортов",
                    "Пётр",
                    "",
                    "imported@example.test",
                    "Import Department",
                    "Import Position",
                    "",
                    "2026-06-01",
                    "",
                    True,
                    "актуален",
                    "трудовой договор",
                    "2026-12-31",
                    120000,
                    104400,
                    "RUB",
                ]
            ]
        )

        response = self.client.post("/employees/import/", {"file": upload})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["preview"]["created"], 1)
        response = self.client.post("/employees/import/", {"action": "confirm"})
        self.assertEqual(response.status_code, 200)
        employee = Employee.objects.get(email="imported@example.test")
        self.assertEqual(employee.department.name, "Import Department")
        self.assertEqual(employee.position.name, "Import Position")
        self.assertTrue(AuditLog.objects.filter(action=AuditLog.ACTION_IMPORT).exists())

    def test_employee_edit_writes_audit_snapshot(self):
        self.client.force_login(self.manager)

        response = self.client.post(
            f"/employees/{self.employee.pk}/edit/",
            {
                "last_name": self.employee.last_name,
                "first_name": "Пётр",
                "middle_name": self.employee.middle_name,
                "email": self.employee.email,
                "department": str(self.department.pk),
                "position": str(self.position.pk),
                "manager": "",
                "hire_date": "2026-01-10",
                "dismissal_date": "",
                "is_active": "on",
                "ipr_status": Employee.IPR_ACTUAL,
                "contract_type": Employee.CONTRACT_LABOR,
                "contract_end_date": "2026-12-31",
                "salary_gross": "150000",
                "salary_net": "130500",
                "currency": "RUB",
            },
        )

        self.assertEqual(response.status_code, 302)
        log = AuditLog.objects.filter(action=AuditLog.ACTION_UPDATE).latest("created_at")
        self.assertEqual(log.actor, self.manager)
        self.assertEqual(log.before["first_name"], "Иван")
        self.assertEqual(log.after["first_name"], "Пётр")

    def test_hrbp_scope_is_limited_to_assigned_department(self):
        group = Group.objects.create(name="HRBP")
        hrbp = make_user("hrbp", hr_perms=["view_employee"])
        hrbp.groups.add(group)
        UserDepartmentAccess.objects.create(user=hrbp, department=self.department, include_children=True)
        other_department = Department.objects.create(name="Sales")
        Employee.objects.create(
            first_name="Outside",
            last_name="Scope",
            email="outside@example.test",
            department=other_department,
            is_active=True,
            ipr_status=Employee.IPR_NONE,
            currency="RUB",
        )

        self.client.force_login(hrbp)
        response = self.client.get("/employees/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["summary"]["count"], 1)

    def test_position_directory_create_and_delete(self):
        self.client.force_login(self.directory_admin)

        response = self.client.post("/positions/new/", {"name": "Temporary Role"})
        self.assertEqual(response.status_code, 302)
        position = Position.objects.get(name="Temporary Role")

        response = self.client.post(f"/positions/{position.pk}/delete/")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Position.objects.filter(pk=position.pk).exists())

    def test_campaign_record_workflow_status_update(self):
        self.client.force_login(self.directory_admin)
        campaign = Campaign.objects.create(
            season="Тест",
            year=2026,
            status=Campaign.STATUS_OPEN,
            start_date="2026-06-01",
        )
        record = CampaignRecord.objects.create(
            campaign=campaign,
            employee=self.employee,
            status=CampaignRecord.STATUS_ACTIVE,
        )

        response = self.client.post(
            f"/campaigns/{campaign.pk}/records/{record.pk}/status/",
            {"status": CampaignRecord.STATUS_HR_APPROVED},
        )

        self.assertEqual(response.status_code, 302)
        record.refresh_from_db()
        self.assertEqual(record.status, CampaignRecord.STATUS_HR_APPROVED)

    def _build_import_file(self, rows):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(
            [
                "ID",
                "Фамилия",
                "Имя",
                "Отчество",
                "Email",
                "Подразделение",
                "Должность",
                "Email руководителя",
                "Дата найма",
                "Дата увольнения",
                "Активен",
                "Статус ИПР",
                "Тип договора",
                "Дата окончания договора",
                "Gross",
                "Net",
                "Валюта",
            ]
        )
        for row in rows:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return SimpleUploadedFile(
            "employees.xlsx",
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
