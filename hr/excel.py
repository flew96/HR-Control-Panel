from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Contract, Department, Employee, Position


EMPLOYEE_XLSX_COLUMNS = [
    ("id", "ID"),
    ("last_name", "Фамилия"),
    ("first_name", "Имя"),
    ("middle_name", "Отчество"),
    ("email", "Email"),
    ("department", "Подразделение"),
    ("position", "Должность"),
    ("manager_email", "Email руководителя"),
    ("hire_date", "Дата найма"),
    ("dismissal_date", "Дата увольнения"),
    ("is_active", "Активен"),
    ("ipr_status", "Статус ИПР"),
    ("contract_type", "Тип договора"),
    ("contract_end_date", "Дата окончания договора"),
    ("salary_gross", "Gross"),
    ("salary_net", "Net"),
    ("currency", "Валюта"),
]

HEADER_TO_KEY = {
    label.lower(): key
    for key, label in EMPLOYEE_XLSX_COLUMNS
}
HEADER_TO_KEY.update({key.lower(): key for key, _label in EMPLOYEE_XLSX_COLUMNS})

IPR_BY_LABEL = {label.lower(): value for value, label in Employee.IPR_STATUS_CHOICES}
IPR_BY_VALUE = {value: value for value, _label in Employee.IPR_STATUS_CHOICES}
CONTRACT_BY_LABEL = {label.lower(): value for value, label in Employee.CONTRACT_TYPE_CHOICES}
CONTRACT_BY_VALUE = {value: value for value, _label in Employee.CONTRACT_TYPE_CHOICES}


def build_employee_workbook(employees, include_rows=True):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Employees"
    headers = [label for _key, label in EMPLOYEE_XLSX_COLUMNS]
    worksheet.append(headers)

    header_fill = PatternFill("solid", fgColor="EAF1FF")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    if include_rows:
        for employee in employees:
            worksheet.append(_employee_row(employee))

    widths = [10, 20, 18, 18, 28, 26, 26, 28, 14, 14, 12, 18, 22, 22, 14, 14, 10]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def import_employees_from_xlsx(file_obj):
    workbook = load_workbook(filename=file_obj, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        return {"created": 0, "updated": 0, "skipped": 0, "errors": ["Файл пустой."]}

    headers = [_normalize_header(header) for header in raw_headers]
    if "last_name" not in headers or "first_name" not in headers:
        return {
            "created": 0,
            "updated": 0,
            "skipped": 0,
            "errors": ["Нужны колонки Фамилия/last_name и Имя/first_name."],
        }

    result = {"created": 0, "updated": 0, "skipped": 0, "errors": []}

    for row_number, raw_row in enumerate(rows, start=2):
        if not raw_row or not any(_has_value(value) for value in raw_row):
            result["skipped"] += 1
            continue

        row = {
            headers[index]: raw_row[index] if index < len(raw_row) else None
            for index in range(len(headers))
            if headers[index]
        }
        try:
            with transaction.atomic():
                created = _upsert_employee(row)
        except ValueError as exc:
            result["errors"].append(f"Строка {row_number}: {exc}")
            continue

        if created:
            result["created"] += 1
        else:
            result["updated"] += 1

    return result


def _employee_row(employee):
    return [
        employee.pk,
        employee.last_name,
        employee.first_name,
        employee.middle_name,
        employee.email,
        employee.department.name if employee.department else "",
        employee.position.name if employee.position else "",
        employee.manager.email if employee.manager and employee.manager.email else "",
        employee.hire_date,
        employee.dismissal_date,
        employee.is_active,
        employee.get_ipr_status_display(),
        employee.get_contract_type_display(),
        employee.contract_end_date,
        employee.salary_gross,
        employee.salary_net,
        employee.currency,
    ]


def _upsert_employee(row):
    employee = _find_employee(row)
    created = employee is None
    if employee is None:
        employee = Employee()

    employee.last_name = _clean_text(row.get("last_name"))
    employee.first_name = _clean_text(row.get("first_name"))
    if not employee.last_name or not employee.first_name:
        raise ValueError("фамилия и имя обязательны.")

    employee.middle_name = _clean_text(row.get("middle_name"))
    employee.email = _clean_text(row.get("email"))
    employee.department = _get_department(row.get("department"))
    employee.position = _get_position(row.get("position"))
    employee.manager = _get_manager(row.get("manager_email"))
    employee.hire_date = _parse_date(row.get("hire_date"))
    employee.dismissal_date = _parse_date(row.get("dismissal_date"))
    employee.is_active = _parse_bool(row.get("is_active"), default=True)
    employee.ipr_status = _choice_value(
        row.get("ipr_status"),
        IPR_BY_VALUE,
        IPR_BY_LABEL,
        Employee.IPR_NONE,
        "статус ИПР",
    )
    employee.contract_type = _choice_value(
        row.get("contract_type"),
        CONTRACT_BY_VALUE,
        CONTRACT_BY_LABEL,
        Employee.CONTRACT_LABOR,
        "тип договора",
    )
    employee.contract_end_date = _parse_date(row.get("contract_end_date"))
    employee.salary_gross = _parse_decimal(row.get("salary_gross"))
    employee.salary_net = _parse_decimal(row.get("salary_net"))
    employee.currency = _choice_value(
        row.get("currency"),
        {"RUB": "RUB", "EUR": "EUR"},
        {"rub": "RUB", "eur": "EUR"},
        "RUB",
        "валюта",
    )
    employee.save()
    Contract.objects.update_or_create(
        employee=employee,
        defaults={
            "contract_type": employee.contract_type,
            "end_date": employee.contract_end_date,
        },
    )
    return created


def _find_employee(row):
    employee_id = row.get("id")
    if _has_value(employee_id):
        try:
            return Employee.objects.get(pk=int(employee_id))
        except (Employee.DoesNotExist, TypeError, ValueError):
            raise ValueError(f"сотрудник с ID {employee_id} не найден.")

    email = _clean_text(row.get("email"))
    if email:
        return Employee.objects.filter(email__iexact=email).first()
    return None


def _get_department(value):
    name = _clean_text(value)
    if not name:
        return None
    department, _created = Department.objects.get_or_create(name=name)
    return department


def _get_position(value):
    name = _clean_text(value)
    if not name:
        return None
    position, _created = Position.objects.get_or_create(name=name)
    return position


def _get_manager(value):
    email = _clean_text(value)
    if not email:
        return None
    return Employee.objects.filter(email__iexact=email).first()


def _normalize_header(header):
    key = _clean_text(header).lower()
    return HEADER_TO_KEY.get(key, "")


def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _has_value(value):
    return value is not None and str(value).strip() != ""


def _parse_date(value):
    if not _has_value(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    for fmt in ["%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"]:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"некорректная дата: {text}.")


def _parse_bool(value, default=False):
    if not _has_value(value):
        return default
    if isinstance(value, bool):
        return value
    text = _clean_text(value).lower()
    if text in {"1", "true", "yes", "да", "активен", "active"}:
        return True
    if text in {"0", "false", "no", "нет", "уволен", "inactive"}:
        return False
    raise ValueError(f"некорректное булево значение: {value}.")


def _parse_decimal(value):
    if not _has_value(value):
        return None
    try:
        return Decimal(str(value).replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"некорректное число: {value}.") from exc


def _choice_value(value, by_value, by_label, default, label):
    if not _has_value(value):
        return default
    text = _clean_text(value)
    if text in by_value:
        return by_value[text]
    lowered = text.lower()
    if lowered in by_label:
        return by_label[lowered]
    raise ValueError(f"некорректное значение поля {label}: {text}.")

